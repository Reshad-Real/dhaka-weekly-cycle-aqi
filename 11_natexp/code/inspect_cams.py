"""Print the layout of each CAMS-BD workbook: first rows, and every row whose first cell is
not a date (station-block headers and unit rows), with its row number."""
import datetime as dt
import sys

import openpyxl

for path in sys.argv[1:]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("=" * 100)
    print(path.replace("\\", "/").split("/")[-1], "| sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        n = 0
        odd = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            n += 1
            if i < 8:
                print("   r%d" % i, list(row[:24]))
            c0 = row[0] if row else None
            if not isinstance(c0, (dt.datetime, dt.date)) and len(odd) < 120:
                odd.append((i, list(row[:24])))
        print("--- sheet %r rows=%d; non-date first cells:" % (ws.title, n))
        for i, r in odd:
            print("   r%d" % i, r)
    wb.close()
