"""Diagnose the 2012-2021 CAMS-BD workbook: sheet names, row count, the type of the
timestamp cell, and examples of each kind of row, per station."""
import collections
import datetime as dt
import sys

import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    kinds = collections.Counter()
    by_station = collections.defaultdict(collections.Counter)
    examples = {}
    n = 0
    for r in ws.iter_rows(values_only=True):
        n += 1
        st = r[0] if r else None
        t = r[1] if len(r) > 1 else None
        if isinstance(t, (dt.datetime, dt.date)):
            k = "datetime"
        elif t is None:
            k = "none"
        elif isinstance(t, str):
            k = "str"
        elif isinstance(t, (int, float)):
            k = "number"
        else:
            k = type(t).__name__
        kinds[k] += 1
        by_station[str(st)][k] += 1
        key = (str(st), k)
        if key not in examples:
            examples[key] = list(r[:12])
    print("sheet %r: %d rows; timestamp kinds %s" % (ws.title, n, dict(kinds)))
    for s, c in by_station.items():
        print("   %-28s %s" % (s[:28], dict(c)))
    print("examples:")
    for (s, k), r in list(examples.items())[:60]:
        print("   [%s|%s]" % (s[:20], k), r)
wb.close()
